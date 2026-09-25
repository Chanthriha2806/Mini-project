import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_review2_presentation():
    template_path = r"E:\Mini Project Documents\REVIEW 1 - AI INFRASTRUCTURE MONITORING DASHBOARD .pptx"
    out_path_1 = r"e:\Mini_Project_Review_2_AI_IT_Infra_Dashboard.pptx"
    out_path_2 = r"e:\Mini_Project_Review_2_AI_IT_Infra_Dashboard_Final.pptx"
    out_path_3 = r"E:\Mini Project Documents\REVIEW 2 - AI INFRASTRUCTURE MONITORING DASHBOARD.pptx"
    dfd_dir = r"e:\Mini project 2\dfd_images"

    if not os.path.exists(template_path):
        print(f"Error: Template not found at {template_path}")
        return

    prs = Presentation(template_path)
    
    # Clear existing slides to rebuild
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]

    # Theme Colors
    DARK_NAVY = RGBColor(15, 23, 42)
    BLUE_ACCENT = RGBColor(30, 64, 175)
    TEXT_DARK = RGBColor(30, 41, 59)
    TEXT_WHITE = RGBColor(255, 255, 255)
    CYAN_HEADER = RGBColor(14, 165, 233)
    BOX_FILL = RGBColor(255, 255, 255)
    STORE_FILL = RGBColor(241, 245, 249)

    def setup_slide_header_footer(slide, title_text):
        tb_title = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(10.5), Inches(0.9))
        tf = tb_title.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text.upper()
        p.font.size = Pt(22)
        p.font.bold = True
        p.font.color.rgb = BLUE_ACCENT

        tb_footer = slide.shapes.add_textbox(Inches(0.3), Inches(7.0), Inches(8.0), Inches(0.4))
        tf_f = tb_footer.text_frame
        p_f = tf_f.paragraphs[0]
        p_f.text = "AI INFRASTRUCTURE MONITORING DASHBOARD"
        p_f.font.size = Pt(10)
        p_f.font.bold = True
        p_f.font.color.rgb = RGBColor(100, 116, 139)

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[0])
    title_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(0.8), Inches(11.33), Inches(5.8))
    title_card.fill.solid()
    title_card.fill.fore_color.rgb = DARK_NAVY
    title_card.line.color.rgb = CYAN_HEADER
    title_card.line.width = Pt(2)

    tb1 = slide1.shapes.add_textbox(Inches(1.3), Inches(1.1), Inches(10.73), Inches(5.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "KGiSL INSTITUTE OF INFORMATION MANAGEMENT"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CYAN_HEADER
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "DEPARTMENT OF COMPUTER SCIENCE | ODD SEMESTER 2026 - 2027"
    p.font.size = Pt(13)
    p.font.color.rgb = RGBColor(148, 163, 184)
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "\nAI-ENABLED IT INFRASTRUCTURE MONITORING DASHBOARD"
    p.font.size = Pt(25)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "MINI PROJECT REVIEW II PRESENTATION\n"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CYAN_HEADER
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "Submitted by: CHANTHRIHA J   |   Reg No: 2538M0070"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = TEXT_WHITE
    p.alignment = PP_ALIGN.CENTER

    p = tf1.add_paragraph()
    p.text = "Under the Guidance of: Mrs. R. GOMATHI (Designation)"
    p.font.size = Pt(14)
    p.font.color.rgb = CYAN_HEADER
    p.alignment = PP_ALIGN.CENTER

    # =========================================================================
    # SLIDE 2: AGENDA / CONTENT
    # =========================================================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide2, "CONTENT - REVIEW II COMPONENTS")

    agenda_items = [
        "1. MODULE DESCRIPTION (Core Subsystems Breakdown)",
        "2. DATA FLOW DIAGRAM (DFD Level 0 Context Diagram)",
        "3. DATA FLOW DIAGRAM (DFD Level 1 Macro Process Flow)",
        "4. DATA FLOW DIAGRAM (DFD Level 2 Detailed Sub-Process Flow)",
        "5. TABLE DESIGN (Relational Database Schemas & Data Types)",
        "6. INPUT DESIGN (Forms, Sliders, JSON Payloads & Validation)",
        "7. OUTPUT DESIGN (Dashboards, Chart.js, Alerts, AI Banners & Reports)",
        "8. SYSTEM REQUIREMENTS & CONCLUSION"
    ]

    top_p = 1.3
    for item in agenda_items:
        card = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(top_p), Inches(11.33), Inches(0.6))
        card.fill.solid()
        card.fill.fore_color.rgb = STORE_FILL
        card.line.color.rgb = BLUE_ACCENT
        
        tb = slide2.shapes.add_textbox(Inches(1.2), Inches(top_p + 0.08), Inches(10.9), Inches(0.45))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = item
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_NAVY
        top_p += 0.7

    # =========================================================================
    # SLIDE 3: COMPONENT 1 - MODULE DESCRIPTION
    # =========================================================================
    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide3, "COMPONENT 1: MODULE DESCRIPTION")

    modules_data = [
        ("1. User Authentication Module", "auth_controller.py & auth_routes.py", "Manages user registration, login authentication using PBKDF2-SHA256 password hashing, Flask-Login session state, and admin vs client role enforcement."),
        ("2. Telemetry Acquisition Module", "system_monitor.py & client_agent.py", "Captures local hardware metrics (CPU %, RAM, Disk, Network speeds) via psutil, or receives remote client JSON payloads via HTTP POST /api/metrics/report."),
        ("3. Threshold Audit & Alert Module", "dashboard_controller.py & mail_helper.py", "Audits telemetry against database threshold limits (settings table); creates ACTIVE alert entries and dispatches async SMTP breach emails."),
        ("4. AI Recommendation Module", "ai_engine.py (RecommendationEngine)", "Rule-based expert diagnostic engine that analyzes multi-metric bottlenecks to output probable root causes and actionable corrective steps."),
        ("5. Reporting & Process Control Module", "reports_controller.py & report_generator.py", "Enables safe PID process termination for administrators, and compiles time-series metrics into PDF (ReportLab) and Excel (openpyxl) audit reports.")
    ]

    top_p = 1.2
    for title, files, desc in modules_data:
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_p), Inches(11.73), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = BOX_FILL
        box.line.color.rgb = BLUE_ACCENT
        box.line.width = Pt(1.5)

        tb = slide3.shapes.add_textbox(Inches(1.0), Inches(top_p + 0.08), Inches(11.33), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"{title}   [{files}]"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(11)
        p1.font.color.rgb = TEXT_DARK
        top_p += 1.1

    # =========================================================================
    # SLIDE 4: COMPONENT 2 - DFD LEVEL 0 (EMBEDDED HIGH-RES PNG DIAGRAM)
    # =========================================================================
    slide4 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide4, "COMPONENT 2: DATA FLOW DIAGRAM (DFD LEVEL 0)")
    
    img0_path = os.path.join(dfd_dir, "dfd_level0.png")
    if os.path.exists(img0_path):
        slide4.shapes.add_picture(img0_path, Inches(0.8), Inches(1.2), width=Inches(11.73))

    # =========================================================================
    # SLIDE 5: COMPONENT 2 - DFD LEVEL 1 (EMBEDDED HIGH-RES PNG DIAGRAM)
    # =========================================================================
    slide5 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide5, "COMPONENT 2: DATA FLOW DIAGRAM (DFD LEVEL 1)")

    img1_path = os.path.join(dfd_dir, "dfd_level1.png")
    if os.path.exists(img1_path):
        slide5.shapes.add_picture(img1_path, Inches(0.8), Inches(1.2), width=Inches(11.73))

    # =========================================================================
    # SLIDE 6: COMPONENT 2 - DFD LEVEL 2 (EMBEDDED HIGH-RES PNG DIAGRAM)
    # =========================================================================
    slide6 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide6, "COMPONENT 2: DATA FLOW DIAGRAM (DFD LEVEL 2)")

    img2_path = os.path.join(dfd_dir, "dfd_level2.png")
    if os.path.exists(img2_path):
        slide6.shapes.add_picture(img2_path, Inches(0.8), Inches(1.2), width=Inches(11.73))

    # =========================================================================
    # SLIDE 7: COMPONENT 3 - TABLE DESIGN (SCHEMAS PART 1)
    # =========================================================================
    slide7 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide7, "COMPONENT 3: TABLE DESIGN (DATABASE SCHEMAS - PART 1)")

    box_u = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.73), Inches(2.5))
    box_u.fill.solid()
    box_u.fill.fore_color.rgb = BOX_FILL
    box_u.line.color.rgb = BLUE_ACCENT

    tb_u = slide7.shapes.add_textbox(Inches(1.0), Inches(1.4), Inches(11.33), Inches(2.3))
    tf_u = tb_u.text_frame
    tf_u.word_wrap = True

    p = tf_u.paragraphs[0]
    p.text = "1. TABLE: `users` (User Authentication & Roles)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p1 = tf_u.add_paragraph()
    p1.text = """
  • id: INT (Primary Key, Auto Increment, Not Null)
  • username: VARCHAR(50) (Unique, Not Null) - Account login username
  • email: VARCHAR(120) (Unique, Not Null) - Registered email address
  • password_hash: VARCHAR(255) (Not Null) - PBKDF2-SHA256 encrypted hash
  • is_admin: BOOLEAN (Default False) - Role flag for administrative privileges
  • created_at: TIMESTAMP (Default Current Timestamp)
    """
    p1.font.size = Pt(11)
    p1.font.color.rgb = TEXT_DARK

    box_s = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.0), Inches(11.73), Inches(2.7))
    box_s.fill.solid()
    box_s.fill.fore_color.rgb = BOX_FILL
    box_s.line.color.rgb = BLUE_ACCENT

    tb_s = slide7.shapes.add_textbox(Inches(1.0), Inches(4.1), Inches(11.33), Inches(2.5))
    tf_s = tb_s.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "2. TABLE: `settings` (User Threshold Configuration)"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p1 = tf_s.add_paragraph()
    p1.text = """
  • id: INT (Primary Key, Auto Increment, Not Null)
  • user_id: INT (Foreign Key -> users.id, Unique, Not Null) - Owner account
  • cpu_threshold: FLOAT (Default 90.0) - CPU usage alarm limit (%)
  • ram_threshold: FLOAT (Default 90.0) - Physical RAM alarm limit (%)
  • disk_threshold: FLOAT (Default 95.0) - Storage space alarm limit (%)
  • refresh_interval: INT (Default 5) - Dashboard live polling interval (Seconds)
  • email_address: VARCHAR(120) - Destination address for SMTP breach alerts
    """
    p1.font.size = Pt(11)
    p1.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 8: COMPONENT 3 - TABLE DESIGN (SCHEMAS PART 2)
    # =========================================================================
    slide8 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide8, "COMPONENT 3: TABLE DESIGN (DATABASE SCHEMAS - PART 2)")

    tables_p2 = [
        ("3. TABLE: `monitoring_history` (Time-Series Metrics Log)",
         "• id: INT (PK, AI) | user_id: INT (FK -> users.id, Nullable)\n• hostname: VARCHAR(100) | cpu_usage: FLOAT | ram_usage: FLOAT | disk_usage: FLOAT\n• upload_speed: FLOAT (KB/s) | download_speed: FLOAT (KB/s) | created_at: TIMESTAMP"),
        ("4. TABLE: `alerts` (Threshold Violation Audit Log)",
         "• id: INT (PK, AI) | metric_type: VARCHAR(20) ('CPU', 'RAM', 'DISK')\n• message: VARCHAR(255) | value: FLOAT | threshold: FLOAT\n• status: VARCHAR(20) (Default 'ACTIVE', 'RESOLVED') | created_at: TIMESTAMP"),
        ("5. TABLE: `reports` (Exported Audit File Registry)",
         "• id: INT (PK, AI) | report_type: VARCHAR(20) ('DAILY', 'WEEKLY', 'MONTHLY')\n• file_path: VARCHAR(255) - File system path to generated PDF/Excel document\n• created_at: TIMESTAMP")
    ]

    top_p = 1.3
    for title, desc in tables_p2:
        box = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_p), Inches(11.73), Inches(1.6))
        box.fill.solid()
        box.fill.fore_color.rgb = BOX_FILL
        box.line.color.rgb = BLUE_ACCENT

        tb = slide8.shapes.add_textbox(Inches(1.0), Inches(top_p + 0.1), Inches(11.33), Inches(1.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(11)
        p1.font.color.rgb = TEXT_DARK
        top_p += 1.8

    # =========================================================================
    # SLIDE 9: COMPONENT 4 - INPUT DESIGN
    # =========================================================================
    slide9 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide9, "COMPONENT 4: INPUT DESIGN (FORMS & CONTROLS)")

    inputs_list = [
        ("User Authentication Input Forms (`login.html` & `register.html`)",
         "Controls: Textbox (Username/Email), Password (Password field).\nValidation: Client-side required field enforcement, email format regex, server-side PBKDF2 hash verification."),
        ("Threshold Configuration Input Form (`settings.html`)",
         "Controls: HTML5 Range Sliders (1-100% for CPU, RAM, Disk limits), Number Input (Refresh rate), Textbox (Alert Email).\nValidation: Float range validation [1.0 - 100.0], valid recipient email verification."),
        ("Remote Client Telemetry JSON Payload (`POST /api/metrics/report`)", "Controls: HTTP REST API JSON POST endpoint from client_agent.py.\nValidation: JSON schema content-type check, mandatory key validation for cpu, memory, disk, network, system."),
        ("AI Diagnostic Assistant Chat Query Input (`dashboard.html`)", "Controls: Natural-language text input box & Send button.\nValidation: Text sanitization, SQL injection prevention, non-empty check."),
        ("Process Termination Action Input (`POST /api/process/kill`)", "Controls: PID Action Button in Top Processes Table.\nValidation: Admin session verification, protection check against core OS system PIDs.")
    ]

    top_p = 1.2
    for title, desc in inputs_list:
        box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_p), Inches(11.73), Inches(1.0))
        box.fill.solid()
        box.fill.fore_color.rgb = BOX_FILL
        box.line.color.rgb = BLUE_ACCENT

        tb = slide9.shapes.add_textbox(Inches(1.0), Inches(top_p + 0.08), Inches(11.33), Inches(0.85))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"• {title}"
        p0.font.size = Pt(13)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(11)
        p1.font.color.rgb = TEXT_DARK
        top_p += 1.1

    # =========================================================================
    # SLIDE 10: COMPONENT 5 - OUTPUT DESIGN (DASHBOARD & ALERTS)
    # =========================================================================
    slide10 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide10, "COMPONENT 5: OUTPUT DESIGN (DASHBOARD & ALERTS)")

    outputs_part1 = [
        ("Real-Time Metric Cards & Circular Progress Badges", "Renders percentage badges and circular progress rings for CPU utilization, RAM usage, Disk storage, and Network upload/download speeds (refreshed every 5 seconds)."),
        ("Interactive Dynamic Multi-Line Chart.js Graphs", "Visualizes historical time-series metric trends over time. Features real-time line smoothing and live push updates via Fetch API polling."),
        ("Active Alert Notification Feed Panel", "Color-coded red/warning notification banners triggered automatically when hardware metrics breach configured thresholds.")
    ]

    top_p = 1.4
    for title, desc in outputs_part1:
        box = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_p), Inches(11.73), Inches(1.5))
        box.fill.solid()
        box.fill.fore_color.rgb = BOX_FILL
        box.line.color.rgb = BLUE_ACCENT

        tb = slide10.shapes.add_textbox(Inches(1.0), Inches(top_p + 0.1), Inches(11.33), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"• {title}"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(12)
        p1.font.color.rgb = TEXT_DARK
        top_p += 1.7

    # =========================================================================
    # SLIDE 11: COMPONENT 5 - OUTPUT DESIGN (AI DIAGNOSTICS & REPORTS)
    # =========================================================================
    slide11 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide11, "COMPONENT 5: OUTPUT DESIGN (AI DIAGNOSTICS & REPORTS)")

    outputs_part2 = [
        ("AI Diagnostic Recommendation Banners & Chat Modal", "Displays structured troubleshooting causes (e.g. memory leak, runaway thread) and step-by-step resolution advice for administrators."),
        ("Top Processes Table & One-Click Admin Controls", "Displays top 10 CPU/RAM consuming processes with PID, username, CPU %, and memory % with single-click PID termination buttons."),
        ("Exported Infrastructure Audit Reports (PDF & Excel)", "Compiles historical time-series statistics into downloadable PDF diagnostic reports (ReportLab Flowables) and Excel workbooks (openpyxl).")
    ]

    top_p = 1.4
    for title, desc in outputs_part2:
        box = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(top_p), Inches(11.73), Inches(1.5))
        box.fill.solid()
        box.fill.fore_color.rgb = BOX_FILL
        box.line.color.rgb = BLUE_ACCENT

        tb = slide11.shapes.add_textbox(Inches(1.0), Inches(top_p + 0.1), Inches(11.33), Inches(1.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p0 = tf.paragraphs[0]
        p0.text = f"• {title}"
        p0.font.size = Pt(14)
        p0.font.bold = True
        p0.font.color.rgb = BLUE_ACCENT

        p1 = tf.add_paragraph()
        p1.text = desc
        p1.font.size = Pt(12)
        p1.font.color.rgb = TEXT_DARK
        top_p += 1.7

    # =========================================================================
    # SLIDE 12: SYSTEM REQUIREMENTS
    # =========================================================================
    slide12 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide12, "SYSTEM REQUIREMENTS & TECHNICAL SPECIFICATIONS")

    card_h = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(5.7), Inches(5.3))
    card_h.fill.solid()
    card_h.fill.fore_color.rgb = BOX_FILL
    card_h.line.color.rgb = BLUE_ACCENT

    tb_h = slide12.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(5.3), Inches(4.9))
    tf_h = tb_h.text_frame
    tf_h.word_wrap = True
    p = tf_h.paragraphs[0]
    p.text = "HARDWARE SPECIFICATIONS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p1 = tf_h.add_paragraph()
    p1.text = """
  • Processor: Intel Core i3 / i5 or equivalent (Minimum 2.0 GHz)
  • RAM: Minimum 4 GB (Recommended 8 GB)
  • Disk Storage: Minimum 10 GB free space for database logs
  • Network Interface: 100/1000 Mbps Ethernet / Wi-Fi
  • Architecture: 64-bit x86/x64 Architecture
    """
    p1.font.size = Pt(12)
    p1.font.color.rgb = TEXT_DARK

    card_sw = slide12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.3), Inches(5.7), Inches(5.3))
    card_sw.fill.solid()
    card_sw.fill.fore_color.rgb = BOX_FILL
    card_sw.line.color.rgb = BLUE_ACCENT

    tb_sw = slide12.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.3), Inches(4.9))
    tf_sw = tb_sw.text_frame
    tf_sw.word_wrap = True
    p = tf_sw.paragraphs[0]
    p.text = "SOFTWARE SPECIFICATIONS"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    p1 = tf_sw.add_paragraph()
    p1.text = """
  • Operating System: Windows 10 / 11 (64-bit)
  • Frontend: HTML5, CSS3, Bootstrap 5, JavaScript, Chart.js
  • Backend: Python 3.12, Flask, Flask-SQLAlchemy, Flask-Login
  • Hardware Telemetry Sensor: psutil library
  • Relational Database: MySQL / SQLite (it_monitor_db)
  • Report Generation: ReportLab (PDF) & openpyxl (Excel)
    """
    p1.font.size = Pt(12)
    p1.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 13: CONCLUSION
    # =========================================================================
    slide13 = prs.slides.add_slide(prs.slide_layouts[1])
    setup_slide_header_footer(slide13, "CONCLUSION & REVIEW II SUMMARY")

    box_c = slide13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.3), Inches(11.73), Inches(5.3))
    box_c.fill.solid()
    box_c.fill.fore_color.rgb = STORE_FILL
    box_c.line.color.rgb = BLUE_ACCENT

    tb_c = slide13.shapes.add_textbox(Inches(1.1), Inches(1.5), Inches(11.13), Inches(4.9))
    tf_c = tb_c.text_frame
    tf_c.word_wrap = True

    p = tf_c.paragraphs[0]
    p.text = "REVIEW II VERIFICATION SUMMARY"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = BLUE_ACCENT

    bullets = [
        "All 5 Mandatory Review II Components Completed: Module Description, DFD Level 0, DFD Level 1, DFD Level 2, Table Design, Input Design, and Output Design.",
        "Centralized Infrastructure Monitoring: Enables real-time tracking of local and remote network nodes via psutil and client_agent.py payloads.",
        "Proactive AI Troubleshooting: Zero-latency rule-based expert recommendation engine formulates root causes and corrective resolution steps.",
        "Automated Threshold Auditing: Triggers instant SMTP email breach alerts and logs ACTIVE violation records in MySQL database.",
        "Compliance Reporting & Administration: Provides single-click export of PDF/Excel performance summaries and safe PID process termination."
    ]

    for b in bullets:
        p = tf_c.add_paragraph()
        p.text = f"\n✔ {b}"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = DARK_NAVY

    prs.save(out_path_1)
    prs.save(out_path_2)
    prs.save(out_path_3)
    print("Successfully built Review II PowerPoint with embedded DFD diagrams!")

if __name__ == '__main__':
    build_review2_presentation()
