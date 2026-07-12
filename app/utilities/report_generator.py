import os
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

class ReportGenerator:
    """Utility class to compile systems metrics statistics reports in PDF and Excel formats."""

    @staticmethod
    def build_pdf(stats, alerts, report_type, filepath):
        """Compiles a professional diagnostic report PDF using ReportLab Flowables."""
        # Ensure directories exist
        os.makedirs(os.path.dirname(filepath), exist_ok=True)
        
        doc = SimpleDocTemplate(filepath, pagesize=letter,
                                rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
        story = []
        
        # Styles
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'ReportTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=24,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=15
        )
        subtitle_style = ParagraphStyle(
            'ReportSub',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#4b5563'),
            spaceAfter=25
        )
        h2_style = ParagraphStyle(
            'SectionHeader',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=14,
            textColor=colors.HexColor('#3b82f6'),
            spaceBefore=15,
            spaceAfter=10
        )
        body_style = ParagraphStyle(
            'BodyTextCustom',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=10,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=10
        )

        # Header Block
        story.append(Paragraph(f"Infrastructure Monitoring Report - {report_type.upper()}", title_style))
        story.append(Paragraph(f"Generated on: {stats.get('generated_at')} | Audit Range: Last {stats.get('range_days')} Days", subtitle_style))
        story.append(Spacer(1, 10))

        # Overview Stats Table
        story.append(Paragraph("System Telemetry Averages", h2_style))
        story.append(Paragraph("Mean and peak resource utilization indexes calculated from historical database timeseries:", body_style))
        
        overview_data = [
            ["Metric Parameter", "Average Value", "Peak Value"],
            ["CPU Core Utilization", f"{stats.get('cpu_avg')}%", f"{stats.get('cpu_max')}%"],
            ["Physical RAM Utilization", f"{stats.get('ram_avg')}%", f"{stats.get('ram_max')}%"],
            ["Disk Volume Utilization", f"{stats.get('disk_avg')}%", f"{stats.get('disk_max')}%"],
            ["Network Upload Speed", f"{stats.get('upload_avg')} KB/s", f"{stats.get('upload_max')} KB/s"],
            ["Network Download Speed", f"{stats.get('download_avg')} KB/s", f"{stats.get('download_max')} KB/s"],
        ]
        
        overview_table = Table(overview_data, colWidths=[200, 160, 160])
        overview_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f2937')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 10),
            ('BOTTOMPADDING', (0, 0), (-1, 0), 8),
            ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f9fafb')),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
            ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 1), (-1, -1), 9),
            ('PADDING', (0, 0), (-1, -1), 6),
        ]))
        story.append(overview_table)
        story.append(Spacer(1, 20))

        # Alerts Feed Section
        story.append(Paragraph("Database Alerts Activity Feed", h2_style))
        story.append(Paragraph(f"Total triggered alarms during audit cycle: {len(alerts)} items.", body_style))

        if alerts:
            alerts_data = [["Timestamp", "Metric", "Threshold Status", "Message", "Breach Value"]]
            for alt in alerts:
                alerts_data.append([
                    alt.get('timestamp'),
                    alt.get('metric_type'),
                    alt.get('status'),
                    alt.get('message'),
                    f"{alt.get('value')}%"
                ])
            
            alerts_table = Table(alerts_data, colWidths=[110, 60, 60, 220, 70])
            alerts_table.setStyle(TableStyle([
                ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1f2937')),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
                ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
                ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                ('FONTSIZE', (0, 0), (-1, 0), 9),
                ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e7eb')),
                ('FONTNAME', (0, 1), (-1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 1), (-1, -1), 8),
                ('PADDING', (0, 0), (-1, -1), 5),
            ]))
            story.append(alerts_table)
        else:
            story.append(Paragraph("Everything looks good! No threshold violations occurred during this reporting timeframe.", body_style))

        doc.build(story)
        return True

    @staticmethod
    def build_excel(stats, alerts, report_type, filepath):
        """Compiles a professional multi-sheet spreadsheet report using openpyxl."""
        # Ensure directories exist
        os.makedirs(os.path.dirname(filepath), exist_ok=True)

        wb = Workbook()
        
        # Styles tokens
        header_font = Font(name='Arial', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1F2937', end_color='1F2937', fill_type='solid')
        title_font = Font(name='Arial', size=14, bold=True, color='1F2937')
        section_font = Font(name='Arial', size=11, bold=True, color='3B82f6')
        bold_font = Font(name='Arial', size=10, bold=True)
        regular_font = Font(name='Arial', size=10)
        border_side = Side(border_style='thin', color='E5E7EB')
        thin_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')

        # 1. Summary Sheet Setup
        ws_summary = wb.active
        ws_summary.title = "Summary Diagnostics"
        ws_summary.views.sheetView[0].showGridLines = True

        ws_summary.cell(row=2, column=2, value=f"IT Infrastructure Monitoring Report - {report_type.upper()}").font = title_font
        ws_summary.cell(row=3, column=2, value=f"Generated on: {stats.get('generated_at')} | Range: Last {stats.get('range_days')} Days").font = Font(name='Arial', size=9, italic=True)

        ws_summary.cell(row=5, column=2, value="Telemetry metrics summary:").font = section_font
        
        # Columns
        headers = ["Metric Parameter", "Average Utilization", "Peak Utilization"]
        for col_idx, h in enumerate(headers, start=2):
            cell = ws_summary.cell(row=6, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        overview_data = [
            ["CPU Utilization", f"{stats.get('cpu_avg')}%", f"{stats.get('cpu_max')}%"],
            ["Physical RAM Allocation", f"{stats.get('ram_avg')}%", f"{stats.get('ram_max')}%"],
            ["Disk Volume Storage", f"{stats.get('disk_avg')}%", f"{stats.get('disk_max')}%"],
            ["Network Upload speed", f"{stats.get('upload_avg')} KB/s", f"{stats.get('upload_max')} KB/s"],
            ["Network Download speed", f"{stats.get('download_avg')} KB/s", f"{stats.get('download_max')} KB/s"],
        ]

        for r_idx, row_data in enumerate(overview_data, start=7):
            for c_idx, val in enumerate(row_data, start=2):
                cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = left_align

        # Auto-size columns summary
        for col in ws_summary.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws_summary.column_dimensions[col_letter].width = max(max_len + 3, 12)

        # 2. Alerts Sheet Setup
        ws_alerts = wb.create_sheet(title="Triggered Alerts")
        ws_alerts.views.sheetView[0].showGridLines = True

        ws_alerts.cell(row=2, column=2, value="Historical Alerts Activity Log").font = title_font
        ws_alerts.cell(row=3, column=2, value=f"Total breached alarms registered: {len(alerts)}").font = Font(name='Arial', size=9, italic=True)

        alert_headers = ["Timestamp", "Metric Class", "Trigger Status", "Alarms Message", "Value at Breach", "Configured Limit"]
        for col_idx, h in enumerate(alert_headers, start=2):
            cell = ws_alerts.cell(row=5, column=col_idx, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = thin_border

        for r_idx, alt in enumerate(alerts, start=6):
            data_row = [
                alt.get('timestamp'),
                alt.get('metric_type'),
                alt.get('status'),
                alt.get('message'),
                f"{alt.get('value')}%",
                f"{alt.get('threshold')}%"
            ]
            for c_idx, val in enumerate(data_row, start=2):
                cell = ws_alerts.cell(row=r_idx, column=c_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = left_align

        # Auto-size columns alerts
        for col in ws_alerts.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = col[0].column_letter
            ws_alerts.column_dimensions[col_letter].width = max(max_len + 3, 12)

        wb.save(filepath)
        return True
